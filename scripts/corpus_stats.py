"""corpus_stats.py — 对一批优秀论文 PDF 计算文风/结构指标，生成 prose_lint 的阈值文件。

用法：
    python scripts/corpus_stats.py <pdf 目录> [--docx 本稿.docx] [--out-json .agents/skills/_references/corpus_stats.json]
                                   [--out-md .agents/skills/_references/corpus_stats.md]

依赖：pymupdf（读 PDF）；--docx 时需要 python-docx。仅在维护语料阈值时运行，输出项目不需要。
统计范围为正文：从"摘要"起到最后一个行首"参考文献"止（排除附录代码，代码里的括号会污染统计）。
仓库只保存统计结果，不保存 PDF。
"""
from __future__ import annotations

import argparse
import json
import re
import statistics as st
import sys
from pathlib import Path

CONNECTIVES = ["首先", "其次", "再次", "然后", "随后", "最后", "从而", "因此", "此外", "进而", "由此", "可以看出",
               "由此可见", "换言之", "一方面", "另一方面", "考虑到", "为此", "于是"]
AI_PHRASES = ["赋能", "需要指出的是", "全方位", "有力支撑", "不难发现", "多维度", "值得一提的是", "众所周知",
              "不言而喻", "无疑", "毫无疑问", "极大地", "极具", "深远影响"]
PROBLEM_HEAD = re.compile(r"问题\s*[一二三四五六七八九十0-9]+")


def zh(s: str) -> int:
    return sum(1 for c in s if "\u4e00" <= c <= "\u9fff")


def cut_body(full: str) -> str:
    start = 0
    m = re.search(r"摘\s*要", full)
    if m:
        start = m.start()
    ends = [m.start() for m in re.finditer(r"(?m)^\s*(参考文献|参 考 文 献|References)\s*$", full)]
    end = ends[-1] if ends else len(full)
    if end <= start:
        end = len(full)
    return full[start:end]


def metrics(body: str) -> dict:
    n = zh(body)
    k = max(n, 1) / 1000.0
    sents = [s for s in re.split(r"[。！？]", body) if zh(s) >= 4]
    slen = [zh(s) for s in sents]
    ai = {p: body.count(p) for p in AI_PHRASES}
    return {
        "zh_chars": n,
        "paren_full_per_k": round(body.count("（") / k, 2),
        "paren_half_zh_per_k": round(len(re.findall(r"\([^)]*[\u4e00-\u9fff][^)]*\)", body)) / k, 2),
        "quote_per_k": round(body.count("“") / k, 2),
        "quote_mismatch": len(re.findall(r"“[^”“]{1,40}“", body)),
        "dash_per_k": round(len(re.findall(r"—+", body)) / k, 2),
        "dash_chain_per_k": round(len(re.findall(r"(?:[^—\n，。]{1,25}—){2,}[^—\n，。]{1,25}", body)) / k, 2),
        "enum_inline_per_k": round(len(re.findall(r"[（(]\s*[1-9①-⑩]\s*[)）]", body)) / k, 2),
        "en_abbr_per_k": round(len(re.findall(r"\b[A-Z]{2,}\b", body)) / k, 2),
        "connective_per_k": round(sum(body.count(c) for c in CONNECTIVES) / k, 2),
        "ai_phrase_per_k": round(sum(ai.values()) / k, 2),
        "ai_hits": {p: c for p, c in ai.items() if c},
        "sent_median": st.median(slen) if slen else 0,
        "sent_p90": sorted(slen)[int(len(slen) * 0.9)] if slen else 0,
        "long_sent_ratio": round(sum(1 for x in slen if x > 80) / max(len(slen), 1), 3),
    }


def toc_metrics(toc: list[tuple[int, str]]) -> dict:
    if not toc:
        return {}
    lv = [l for l, _ in toc]
    return {"h1": lv.count(1), "h2": lv.count(2), "h3": lv.count(3), "h4plus": sum(1 for l in lv if l >= 4),
            "h3_per_h2": round(lv.count(3) / max(lv.count(2), 1), 2),
            "paren_title_ratio": round(sum(1 for _, t in toc if "（" in t or "(" in t) / len(toc), 2)}


def pdf_stats(path: Path) -> tuple[dict, list[int]]:
    import pymupdf
    doc = pymupdf.open(path)
    pages = [p.get_text() for p in doc]
    toc = doc.get_toc()
    row = {"name": path.stem, "pages": doc.page_count, **metrics(cut_body("\n".join(pages))),
           **toc_metrics([(l, t) for l, t, _ in toc])}
    # 各"问题 N"一级章节的汉字数
    l1 = [(t, p) for l, t, p in toc if l == 1]
    chapters = []
    for (t, p), (_, p2) in zip(l1, l1[1:] + [("", len(pages) + 1)]):
        if PROBLEM_HEAD.search(t) and "重述" not in t and not re.search(r"^\S*\s*问题分析$", t):
            chapters.append(zh("".join(pages[p - 1:p2 - 1])))
    row["problem_chapters"] = chapters
    return row, chapters


def docx_stats(path: Path) -> dict:
    from docx import Document
    d = Document(str(path))
    lines, toc = [], []
    for p in d.paragraphs:
        s = p.style.name if p.style is not None else ""
        m = re.match(r"Heading (\d)", s)
        if m:
            toc.append((int(m.group(1)), p.text))
        lines.append(p.text)
    for t in d.tables:
        for row in t.rows:
            lines.append(" ".join(c.text for c in row.cells))
    return {"name": path.stem, "pages": None, **metrics(cut_body("\n".join(lines))), **toc_metrics(toc)}


KEYS = ["zh_chars", "paren_full_per_k", "paren_half_zh_per_k", "quote_per_k", "dash_per_k", "dash_chain_per_k",
        "enum_inline_per_k", "en_abbr_per_k", "connective_per_k", "ai_phrase_per_k", "sent_median", "sent_p90",
        "long_sent_ratio", "h1", "h2", "h3", "h3_per_h2", "paren_title_ratio"]


def quantiles(vals: list[float]) -> dict:
    q = st.quantiles(vals, n=4) if len(vals) >= 2 else [vals[0]] * 3
    return {"p25": round(q[0], 2), "p50": round(q[1], 2), "p75": round(q[2], 2), "max": max(vals), "min": min(vals), "n": len(vals)}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pdf_dir")
    ap.add_argument("--docx", default=None)
    ap.add_argument("--out-json", default=None)
    ap.add_argument("--out-md", default=None)
    a = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    rows, chapters = [], []
    for f in sorted(Path(a.pdf_dir).glob("*.pdf")):
        r, ch = pdf_stats(f)
        rows.append(r)
        chapters += ch
        print(f"[corpus] {f.name}: {r['pages']} 页，正文 {r['zh_chars']} 字，问题章节 {ch}")
    ours = docx_stats(Path(a.docx)) if a.docx else None
    summary = {k: quantiles([r[k] for r in rows if isinstance(r.get(k), (int, float))]) for k in KEYS}
    summary = {k: v for k, v in summary.items() if v["n"]}
    summary["problem_chapter_zh"] = quantiles(chapters) if chapters else None
    ai_total: dict[str, int] = {}
    for r in rows:
        for p, c in r["ai_hits"].items():
            ai_total[p] = ai_total.get(p, 0) + c
    out = {"n_papers": len(rows), "rows": rows, "ours": ours, "summary": summary, "ai_phrase_corpus_hits": ai_total}

    md = ["# 优秀论文语料文风统计", "", f"n = {len(rows)} 篇；正文 = 摘要起、参考文献止；密度按每千汉字。由 `scripts/corpus_stats.py` 生成，勿手改。", "",
          "| 指标 | P25 | P50 | P75 | max |" + (" 本稿 |" if ours else ""), "| --- | ---: | ---: | ---: | ---: |" + (" ---: |" if ours else "")]
    for k in KEYS:
        if k not in summary:
            continue
        s = summary[k]
        md.append(f"| {k} | {s['p25']} | {s['p50']} | {s['p75']} | {s['max']} |" + (f" {ours.get(k, '-')} |" if ours else ""))
    if summary["problem_chapter_zh"]:
        s = summary["problem_chapter_zh"]
        md.append(f"| problem_chapter_zh（{s['n']} 章） | {s['p25']} | {s['p50']} | {s['p75']} | {s['max']} |" + (" |" if ours else ""))
    md += ["", "## 语料中 AI 高频措辞的出现次数（全部论文合计）", "",
           "、".join(f"{p}: {ai_total.get(p, 0)}" for p in AI_PHRASES), "",
           "## 各篇明细", "", "| 论文 | 页 | 汉字 | （/千字 | “/千字 | —/千字 | 缩写/千字 | 连接词/千字 | H1/H2/H3 | 问题章节字数 |",
           "| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | --- | --- |"]
    for r in rows:
        md.append(f"| {r['name']} | {r['pages']} | {r['zh_chars']} | {r['paren_full_per_k']} | {r['quote_per_k']} | {r['dash_per_k']} | "
                  f"{r['en_abbr_per_k']} | {r['connective_per_k']} | {r.get('h1', '-')}/{r.get('h2', '-')}/{r.get('h3', '-')} | {r['problem_chapters']} |")
    text = "\n".join(md) + "\n"
    print(text)
    if a.out_json:
        Path(a.out_json).write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    if a.out_md:
        Path(a.out_md).write_text(text, encoding="utf-8")


if __name__ == "__main__":
    main()
